"""สร้างสไลด์นำเสนอ Mumbai_KT_EDA_Presentation.pptx
สไตล์ "lab-notebook" โทนอุ่น (อ้างอิงดีไซน์จาก hw3/Decoding_Wheat_Seed_Geometry.pptx)
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.lang import MSO_LANGUAGE_ID
from pptx.oxml.ns import qn
from pathlib import Path

CH = Path("charts")

# ---------- palette (โทนอุ่นแบบ lab-notebook — เข้มขึ้น อิ่มขึ้นกว่าเดิม) ----------
CREAM   = RGBColor(0xF6, 0xF1, 0xE6)   # พื้นครีม
PAPER   = RGBColor(0xFF, 0xFD, 0xF8)   # การ์ดขาวนวล
INK     = RGBColor(0x2B, 0x26, 0x20)   # หมึกเข้ม
BODY    = RGBColor(0x47, 0x41, 0x38)
MUTED   = RGBColor(0x8C, 0x82, 0x73)
LINE    = RGBColor(0xDD, 0xD2, 0xBE)

BLUE    = RGBColor(0x3E, 0x63, 0x92)   # slate เข้ม
RUST    = RGBColor(0xB4, 0x55, 0x39)   # terracotta อิ่ม
SAGE    = RGBColor(0x63, 0x8B, 0x52)   # green เข้ม
GOLD    = RGBColor(0xC7, 0x97, 0x3B)   # เหลืองอำพัน (accent เสริม)

BLUE_T  = RGBColor(0xE5, 0xEB, 0xF2)   # tint อ่อน
RUST_T  = RGBColor(0xF3, 0xE3, 0xDA)
SAGE_T  = RGBColor(0xE7, 0xEE, 0xDF)
GOLD_T  = RGBColor(0xF4, 0xE9, 0xD1)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)

INK_T = RGBColor(0xEC, 0xE5, 0xD6)
TINT = {str(BLUE): BLUE_T, str(RUST): RUST_T, str(SAGE): SAGE_T,
        str(GOLD): GOLD_T, str(INK): INK_T}


def tint(accent):
    return TINT.get(str(accent), INK_T)


HEAD = "Sarabun"
BODYF = "Sarabun"

AUTHOR = "ธนเสริฐ ตั้งตรงสกุลดี"
SID    = "6805007181"
COURSE = "COS3302 · Data Science — สอบกลางภาค"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


# ============================================================ helpers
def _set_typefaces(rPr, name):
    """ตั้งเฉพาะ <a:latin> — LibreOffice headless เครื่องนี้ถ้าใส่ cs/ea ด้วย อักษรไทยจะหาย"""
    el = rPr.find(qn("a:latin"))
    if el is None:
        el = rPr.makeelement(qn("a:latin"), {})
        rPr.append(el)
    el.set("typeface", name)


def _pstyle(p, name=BODYF):
    try:
        _set_typefaces(p._p.get_or_add_pPr().get_or_add_defRPr(), name)
    except Exception:
        pass


def _rfont(run, size, *, bold=False, color=BODY, italic=False, name=BODYF):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color


def _font(run_or_para, size, *, bold=False, color=BODY, name=BODYF, italic=False):
    if hasattr(run_or_para, "_p"):
        _pstyle(run_or_para, name)
        f = run_or_para.font
    else:
        f = run_or_para.font if hasattr(run_or_para, "font") else run_or_para
        try:
            _set_typefaces(f._rPr, name)
        except Exception:
            pass
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color


def _rect(slide, x, y, w, h, fill, line=None, line_w=1.0, round_=False, radius=0.045):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if round_:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def _dot(slide, x, y, d, color):
    o = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    o.fill.solid(); o.fill.fore_color.rgb = color
    o.line.fill.background(); o.shadow.inherit = False
    return o


def _text(slide, x, y, w, h, runs, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          line_spacing=1.0, wrap=True):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        if isinstance(para, str):
            para = [(para, {})]
        items = para
        pname = items[0][1].get("name", BODYF) if items else BODYF
        _pstyle(p, pname)
        for txt, kw in items:
            r = p.add_run(); r.text = txt
            _rfont(r, kw.get("size", 16), bold=kw.get("bold", False),
                   color=kw.get("color", BODY), italic=kw.get("italic", False))
        sp = items[0][1].get("space_after")
        if sp:
            p.space_after = Pt(sp)
    return tb


def base_slide(framed=True, bg=CREAM):
    s = prs.slides.add_slide(BLANK)
    _rect(s, 0, 0, SW, SH, bg)
    if framed:
        _rect(s, Inches(0.30), Inches(0.30), SW - Inches(0.60), SH - Inches(0.60),
              None, line=LINE, line_w=1.2, round_=True, radius=0.03)
    return s


def footer(s):
    _text(s, Inches(0.6), SH - Inches(0.54), Inches(7.0), Inches(0.32),
          [[(COURSE, dict(size=10, color=MUTED))]])
    _text(s, SW - Inches(6.1), SH - Inches(0.54), Inches(5.5), Inches(0.32),
          [[(f"{AUTHOR} · {SID}", dict(size=10, color=MUTED))]], align=PP_ALIGN.RIGHT)


def head(s, title, kicker=None, accent=INK):
    # แถบสีบางด้านบน = เอกลักษณ์ประจำส่วน
    _rect(s, 0, 0, SW, Inches(0.13), accent)
    y = Inches(0.6)
    if kicker:
        pill = _rect(s, Inches(0.72), y, Inches(0.28) + Inches(0.092) * len(kicker),
                     Inches(0.36), tint(accent), round_=True, radius=0.5)
        tfp = pill.text_frame
        tfp.margin_left = tfp.margin_right = Pt(10)
        tfp.margin_top = tfp.margin_bottom = Pt(2)
        tfp.vertical_anchor = MSO_ANCHOR.MIDDLE
        _pstyle(tfp.paragraphs[0], HEAD)
        rr = tfp.paragraphs[0].add_run(); rr.text = kicker.upper()
        _rfont(rr, 11, bold=True, color=accent)
        y = y + Inches(0.5)
    _text(s, Inches(0.72), y, Inches(11.9), Inches(1.0),
          [[(title, dict(size=28, bold=True, color=INK, name=HEAD))]])
    _rect(s, Inches(0.74), y + Inches(0.92), Inches(1.6), Pt(4), accent, round_=True, radius=0.5)
    _dot(s, Inches(2.46), y + Inches(0.87), Inches(0.14), tint(accent))
    return y + Inches(1.2)


# ============================================================ 1 · ปก
s = base_slide(framed=False)
# แถบสีซ้าย 3 ช่วง
for i, c in enumerate([BLUE, RUST, SAGE]):
    _rect(s, 0, int(i * SH / 3), Inches(0.18), int(SH / 3) + 1, c)
# บล็อกสีอ่อนหลังโมทีฟ
_rect(s, Inches(8.55), Inches(0.7), Inches(4.2), Inches(6.1), SAGE_T, round_=True, radius=0.06)
_text(s, Inches(0.95), Inches(1.05), Inches(7.5), Inches(3.0),
      [[("ปัจจัยเสี่ยงของการ", dict(size=41, bold=True, color=INK, name=HEAD))],
       [("ติด KT", dict(size=41, bold=True, color=RUST, name=HEAD)),
        ("  ในนักศึกษา", dict(size=41, bold=True, color=INK, name=HEAD))],
       [("มหาวิทยาลัยมุมไบ", dict(size=41, bold=True, color=INK, name=HEAD))]],
      line_spacing=1.04)
_text(s, Inches(0.97), Inches(3.62), Inches(7.3), Inches(1.5),
      [[("การสำรวจและวิเคราะห์ข้อมูลเบื้องต้น (Exploratory Data Analysis)",
         dict(size=17, color=BODY))],
       [("ค่าสถิติเชิงพรรณนา · ความสัมพันธ์ระหว่างตัวแปร", dict(size=17, color=BODY))]],
      line_spacing=1.25)
# โมทีฟ: แท่งข้อมูล + จุด
bx, by, bw, gp = Inches(9.05), Inches(5.7), Inches(0.62), Inches(0.34)
for i, (c, hgt) in enumerate([(BLUE, 2.3), (RUST, 3.6), (SAGE, 1.7), (GOLD, 2.9)]):
    _rect(s, bx + i * (bw + gp), by - Inches(hgt), bw, Inches(hgt), c, round_=True, radius=0.12)
for i, c in enumerate([SAGE, BLUE, GOLD, SAGE]):
    _dot(s, bx + i * (bw + gp) + Inches(0.19), by - Inches(4.15), Inches(0.24), c)
# accent rule + tagline
_rect(s, Inches(0.97), Inches(5.1), Inches(1.5), Pt(4), RUST, round_=True, radius=0.5)
_text(s, Inches(0.97), Inches(5.32), Inches(7.4), Inches(1.1),
      [[("ชุดข้อมูล 10,000 records × 21 features", dict(size=15, bold=True, color=INK))],
       [("จำแนกปัจจัยที่ทำให้นักศึกษาสอบไม่ผ่านและต้องสอบแก้ (Keep Terms)",
         dict(size=13, color=MUTED))]], line_spacing=1.3)
_rect(s, Inches(0.95), Inches(6.6), SW - Inches(1.9), Pt(1.2), LINE)
_text(s, Inches(0.97), Inches(6.74), Inches(11.6), Inches(0.5),
      [[("วิชา COS3302 — Data Science · สอบกลางภาค    |    ชื่อ–สกุล: ", dict(size=12.5, color=BODY)),
        (AUTHOR, dict(size=12.5, bold=True, color=INK)),
        ("    |    รหัสนักศึกษา: ", dict(size=12.5, color=BODY)),
        (SID, dict(size=12.5, bold=True, color=INK))]])

# ============================================================ 2 · 3 คอลัมน์ concept
s = base_slide(); footer(s)
head(s, "ที่มา บริบท และความสำคัญของชุดข้อมูล", kicker="ส่วนที่ 1 · Overview", accent=BLUE)
cols = [
    (BLUE, "1", "ที่มา (Source)",
     ["ชุดข้อมูล Mumbai University KT Students",
      "นักศึกษาคณะวิศวกรรมศาสตร์ในเครือ University of Mumbai (อินเดีย)",
      "1 แถว = นักศึกษา 1 คน · 10,000 แถว × 21 คอลัมน์"]),
    (RUST, "2", "บริบท: KT คืออะไร",
     ["KT = \"Keep Terms\" = สอบไม่ผ่านรายวิชา แต่ยังเลื่อนชั้นได้",
      "ต้องสอบแก้ (backlog) วิชานั้นภายหลัง",
      "เทียบได้กับการ \"ติด F / ติดค้างวิชา\" ในระบบไทย"]),
    (SAGE, "3", "ความสำคัญ (Why)",
     ["อัตราการติด KT = ตัวชี้วัดความเสี่ยงการคงอยู่และการจบการศึกษา",
      "ตัวแปรครบทั้ง demographic / พฤติกรรม / ระบบสนับสนุน / ผลลัพธ์",
      "รู้ปัจจัยล่วงหน้า → ออกแบบระบบ early-warning ได้"]),
]
cw = Inches(3.92); gap = Inches(0.28); x0 = Inches(0.72); y0 = Inches(2.4)
for i, (c, num, ttl, lines) in enumerate(cols):
    cx = x0 + i * (cw + gap)
    _rect(s, cx, y0, cw, Inches(4.35), tint(c), round_=True, radius=0.05)
    disc = _dot(s, cx + Inches(0.32), y0 + Inches(0.34), Inches(0.72), c)
    tfn = disc.text_frame; tfn.word_wrap = False
    tfn.paragraphs[0].alignment = PP_ALIGN.CENTER
    _pstyle(tfn.paragraphs[0], HEAD)
    rn = tfn.paragraphs[0].add_run(); rn.text = num
    _rfont(rn, 19, bold=True, color=WHITE)
    _text(s, cx + Inches(0.32), y0 + Inches(1.24), cw - Inches(0.6), Inches(0.5),
          [[(ttl, dict(size=17, bold=True, color=c, name=HEAD))]])
    _text(s, cx + Inches(0.32), y0 + Inches(1.86), cw - Inches(0.62), Inches(2.3),
          [[("•  " + ln, dict(size=13, color=BODY))] for ln in lines],
          line_spacing=1.28)

# ============================================================ big-number helper
def bignum_slide(title, kicker, accent, stats):
    s = base_slide(); footer(s)
    head(s, title, kicker=kicker, accent=accent)
    cw = Inches(3.94); x0 = Inches(0.74); y0 = Inches(2.45); ch = Inches(3.9)
    for i, (c, big, lab, sub) in enumerate(stats):
        cx = x0 + i * (cw + Inches(0.2))
        _rect(s, cx, y0, cw, ch, tint(c), round_=True, radius=0.05)
        _rect(s, cx, y0, cw, Pt(5), c, round_=True, radius=0.5)
        _text(s, cx + Inches(0.3), y0 + Inches(0.35), cw - Inches(0.6), Inches(1.5),
              [[(big, dict(size=58, bold=True, color=c, name=HEAD))]])
        _text(s, cx + Inches(0.32), y0 + Inches(1.86), cw - Inches(0.62), Inches(0.6),
              [[(lab, dict(size=16, bold=True, color=INK))]])
        _text(s, cx + Inches(0.32), y0 + Inches(2.5), cw - Inches(0.62), Inches(1.2),
              [[(sub, dict(size=13, color=BODY))]], line_spacing=1.25)
    return s

# ============================================================ 3 · ตารางกลุ่มตัวแปร (+ แถบตัวเลข)
def table_slide(title, headers, rows, kicker, accent, col_w, fsize=13, top=Inches(2.2)):
    s = base_slide(); footer(s)
    head(s, title, kicker=kicker, accent=accent)
    nrow, ncol = len(rows) + 1, len(headers)
    gt = s.shapes.add_table(nrow, ncol, Inches(0.72), top,
                            Inches(11.9), Inches(0.46 * nrow)).table
    try:
        gt.first_row = False; gt.horz_banding = False
    except Exception:
        pass
    for j, wv in enumerate(col_w):
        gt.columns[j].width = Inches(wv)
    for j, h in enumerate(headers):
        c = gt.cell(0, j); c.text = h
        c.fill.solid(); c.fill.fore_color.rgb = accent
        _font(c.text_frame.paragraphs[0], fsize + 1, bold=True, color=WHITE, name=HEAD)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = gt.cell(i, j); c.text = str(val)
            c.fill.solid()
            if j == 0:
                c.fill.fore_color.rgb = tint(accent)
            else:
                c.fill.fore_color.rgb = PAPER if i % 2 else CREAM
            _font(c.text_frame.paragraphs[0], fsize,
                  color=INK if j == 0 else BODY, bold=(j == 0), name=BODYF)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
    return s

_sv = table_slide(
    "โครงสร้างชุดข้อมูล — 21 คอลัมน์ · 6 กลุ่ม",
    ["กลุ่ม", "คอลัมน์"],
    [["รหัส (1)", "Student_ID"],
     ["Demographic (4)", "Age · Gender · Branch · Year_of_Study"],
     ["การเดินทาง (2)", "Commute_Mode · Commute_Time_mins"],
     ["พฤติกรรม / เวลา (4)", "Attendance_% · Study_Hours · Social_Media_Hours · Sleep_Hours"],
     ["ระบบสนับสนุน (5)", "Part_Time_Job · Extracurricular · Peer_Learning · Mentorship · Financial_Status"],
     ["ผลลัพธ์ / เป้าหมาย (5)", "Previous_Semester_SGPA · Stress_Level · Number_of_KTs · Has_KT · Cleared_in_First_Attempt"]],
    kicker="ส่วนที่ 2 · Descriptive", accent=RUST, col_w=[3.0, 8.9], fsize=12.5,
    top=Inches(3.5))
# แถบตัวเลขภาพรวม 3 ก้อน เหนือตาราง
_chipw = Inches(3.86)
for _i, (_c, _n, _l) in enumerate([(BLUE, "10,000", "แถว (rows)"),
                                   (RUST, "21", "คอลัมน์ (columns)"),
                                   (SAGE, "74%", "ติด KT ≥ 1 วิชา")]):
    _cx = Inches(0.72) + _i * (_chipw + Inches(0.16))
    _rect(_sv, _cx, Inches(2.5), _chipw, Inches(0.82), tint(_c), round_=True, radius=0.16)
    _text(_sv, _cx + Inches(0.3), Inches(2.58), _chipw - Inches(0.5), Inches(0.66),
          [[(_n + "   ", dict(size=23, bold=True, color=_c, name=HEAD)),
            (_l, dict(size=12.5, color=BODY))]], anchor=MSO_ANCHOR.MIDDLE)

# ============================================================ 5 · big numbers (คุณภาพข้อมูล)
bignum_slide("คุณภาพข้อมูล", "ส่วนที่ 2 · Data Quality", RUST, [
    (BLUE, "0", "แถวซ้ำ / รหัสซ้ำ", "ข้อมูลสะอาด ไม่ต้อง de-duplicate"),
    (RUST, "~2%", "ค่าหาย ใน 4 คอลัมน์ตัวเลข", "Study / Social_Media / Sleep / SGPA · กระจายแบบสุ่ม (คล้าย MCAR)"),
    (SAGE, "3 → 1", "คอลัมน์เป้าหมายที่ซ้ำซ้อน", "Has_KT ≡ (KTs>0) ≡ ตรงข้าม Cleared → ระวัง target leakage"),
])

# ============================================================ chart slide helper
def chart_slide(title, img, bullets, kicker, accent, img_w=Inches(9.5)):
    """รูปกราฟ (จัดกึ่งกลางในโซนซ้าย) + การ์ดคำอธิบายสีอ่อนด้านขวา ตัวอักษร 16pt"""
    s = base_slide(); footer(s)
    head(s, title, kicker=kicker, accent=accent)
    band_top = Inches(2.22)
    band_h = SH - band_top - Inches(0.72)

    CARD_W = Inches(4.35)
    card_x = SW - Inches(0.55) - CARD_W
    zone_l = Inches(0.5)
    zone_r = card_x - Inches(0.25)
    max_img_w = zone_r - zone_l

    pic = s.shapes.add_picture(str(CH / img), zone_l, band_top,
                               width=min(img_w, max_img_w))
    if pic.height > band_h:
        r = band_h / pic.height
        pic.height = int(pic.height * r); pic.width = int(pic.width * r)
    pic.left = int(zone_l + (max_img_w - pic.width) / 2)
    pic.top = int(band_top + (band_h - pic.height) / 2)

    # การ์ดคำอธิบาย
    _rect(s, card_x, band_top, CARD_W, band_h, tint(accent), round_=True, radius=0.045)
    _rect(s, card_x, band_top, Pt(5), band_h, accent, round_=True, radius=0.5)
    runs = []
    for b in bullets:
        if b.startswith("#"):
            runs.append([(b[1:].strip(), dict(size=18, bold=True, color=accent, name=HEAD,
                                              space_after=3))])
        else:
            runs.append([("•  " + b, dict(size=16, color=BODY, space_after=6))])
    _text(s, card_x + Inches(0.28), band_top + Inches(0.15), CARD_W - Inches(0.5),
          band_h - Inches(0.3), runs, line_spacing=1.24, anchor=MSO_ANCHOR.MIDDLE)
    return s

# ============================================================ 6 · missing chart
chart_slide("แผนที่ค่าหาย (Missing-Value Map)", "01_missing_values.png",
            ["#อ่านกราฟ",
             "ซ้าย: อัตราค่าหายรายคอลัมน์ อยู่ที่ ~2% เท่านั้น",
             "ขวา: จุดค่าหายกระจายทั่วตาราง ไม่เกาะกลุ่ม",
             "#นัยยะ",
             "รูปแบบคล้าย MCAR → เติมด้วย median / model-based ได้",
             "ไม่จำเป็นต้องตัดแถวทิ้ง"],
            kicker="ส่วนที่ 2 · Data Quality", accent=RUST)

# ============================================================ 6 · describe table (รวม skew/kurtosis)
table_slide(
    "สถิติเชิงพรรณนา — ตัวแปรตัวเลขทั้ง 8",
    ["feature", "mean", "std", "min", "median", "max", "skew", "kurtosis"],
    [["Age", "21.02", "2.02", "18", "21", "24", "-0.01", "-1.28"],
     ["Commute_Time_mins", "69.10", "35.11", "10", "67", "180", "0.30", "-0.54"],
     ["Attendance_Percentage", "69.19", "14.64", "20", "69", "100", "-0.13", "-0.23"],
     ["Study_Hours_per_Day", "2.99", "1.47", "0", "3.0", "8.5", "0.10", "-0.33"],
     ["Social_Media_Hours_per_Day", "3.51", "1.49", "0", "3.5", "8.0", "0.07", "-0.23"],
     ["Sleep_Hours_per_Night", "6.52", "1.19", "3.0", "6.5", "10.0", "0.02", "-0.13"],
     ["Previous_Semester_SGPA", "5.53", "0.86", "2.32", "5.52", "9.08", "0.02", "-0.04"],
     ["Number_of_KTs", "1.54", "1.33", "0", "1", "7", "0.79", "0.44"]],
    kicker="ส่วนที่ 2 · Descriptive", accent=RUST,
    col_w=[3.5, 1.25, 1.25, 1.1, 1.35, 1.15, 1.15, 1.45], fsize=12, top=Inches(2.35))
_text(prs.slides[-1], Inches(0.74), Inches(6.35), Inches(11.9), Inches(0.9),
      [[("ทุกตัว |skew| < 0.8 → ค่อนข้างสมมาตร (mean ≈ median)  ·  "
         "kurtosis ติดลบเกือบทั้งหมด → หางบาง / ค่อนข้างแบน",
         dict(size=13, color=MUTED))]], line_spacing=1.2)

# ============================================================ 9 · categorical freq table
table_slide(
    "ตัวแปรจัดกลุ่ม — การกระจายเด่น",
    ["ตัวแปร", "การกระจาย (%)"],
    [["Gender", "Male 59.0 · Female 39.2 · Other 1.8   (ไม่สมดุล)"],
     ["Branch", "Computer Eng. 25 · IT 20 · AI&DS 15 · EXTC 14 · อื่น ๆ 26"],
     ["Year_of_Study", "First / Second / Third / Final ≈ 25% เท่ากัน"],
     ["Commute_Mode", "Local Train 55 · Bus 20 · Two-Wheeler 15 · Walking 5 · Auto/Cab 5"],
     ["Financial_Status", "Middle 55 · Upper Middle 25 · Lower Middle 15 · Wealthy 5"],
     ["Stress_Level", "Low 50 · Medium 32 · High 18"],
     ["Has_KT", "Yes 74.1 · No 25.9   (class imbalance)"]],
    kicker="ส่วนที่ 2 · Descriptive", accent=RUST, col_w=[2.6, 9.3], fsize=13.5)

# ============================================================ 10-12 univariate charts
chart_slide("การแจกแจงตัวแปรตัวเลขทั้ง 8", "02_hist_numeric.png",
            ["#รูปทรงการแจกแจง",
             "เส้นแดง = mean, เส้นเขียว = median",
             "เกือบทุกตัวสมมาตร mean ≈ median",
             "#ข้อสังเกต",
             "Number_of_KTs เบ้ขวาเล็กน้อย (ข้อมูลนับ)",
             "Age เป็นแบบไม่ต่อเนื่อง (18–24 ปี)"],
            kicker="ส่วนที่ 2 · Univariate", accent=RUST)
chart_slide("ตัวแปรเป้าหมาย — Number_of_KTs & Has_KT", "04_target_distribution.png",
            ["#เป้าหมาย",
             "ค่าเฉลี่ย 1.54 วิชา/คน · median 1 · สูงสุด 7",
             "74% ของนักศึกษาติด KT อย่างน้อย 1 วิชา",
             "#นัยยะต่อการสร้างโมเดล",
             "คลาสไม่สมดุล (74 : 26)",
             "ควรใช้ metric เช่น F1 / ROC-AUC แทน accuracy"],
            kicker="ส่วนที่ 2 · Target", accent=RUST)

# ============================================================ 13 · section 3 divider
s = base_slide(framed=False)
for i, c in enumerate([BLUE, RUST, SAGE]):
    _rect(s, 0, int(i * SH / 3), Inches(0.18), int(SH / 3) + 1, c)
# จุดสีตกแต่งมุมขวาบน
for i, c in enumerate([SAGE, RUST, BLUE]):
    _dot(s, SW - Inches(1.7) + i * Inches(0.42), Inches(1.0), Inches(0.26), c)
_text(s, Inches(1.0), Inches(1.5), Inches(9.0), Inches(0.5),
      [[("ส่วนที่ 3  ·  Relationships", dict(size=16, bold=True, color=SAGE, name=HEAD))]])
_text(s, Inches(0.97), Inches(2.0), Inches(11.0), Inches(1.4),
      [[("ความสัมพันธ์ระหว่างตัวแปร", dict(size=42, bold=True, color=INK, name=HEAD))]])
_rect(s, Inches(1.0), Inches(3.3), Inches(1.7), Pt(4), SAGE, round_=True, radius=0.5)
panel_y = Inches(3.75)
_rect(s, Inches(0.95), panel_y, Inches(11.0), Inches(2.7), PAPER, line=LINE, round_=True, radius=0.05)
_rect(s, Inches(0.95), panel_y, Pt(5), Inches(2.7), SAGE, round_=True, radius=0.5)
rows = [("ตัวเลข × ตัวเลข", "Pearson & Spearman correlation · regression"),
        ("ตัวเลข × กลุ่ม", "boxplot / violin · ค่าเฉลี่ยและมัธยฐานแยกกลุ่ม"),
        ("กลุ่ม × กลุ่ม", "crosstab · stacked bar · Cramér's V (ค่าความสัมพันธ์)")]
for i, (a, b) in enumerate(rows):
    yy = panel_y + Inches(0.36) + i * Inches(0.72)
    _dot(s, Inches(1.34), yy + Inches(0.09), Inches(0.17), SAGE)
    _text(s, Inches(1.74), yy, Inches(10.0), Inches(0.6),
          [[(a + "     ", dict(size=17, bold=True, color=INK)),
            (b, dict(size=14, color=MUTED))]])
footer(s)

# ============================================================ 14 · correlation heatmap
chart_slide("ความสัมพันธ์ ตัวเลข × ตัวเลข", "05_correlation_heatmap.png",
            ["#สัมพันธ์กับ Number_of_KTs",
             "SGPA เดิม  r ≈ −0.35  (แรงสุด)",
             "Study_Hours −0.30 · Attendance −0.30",
             "Social_Media +0.18 · Commute_Time +0.11",
             "#โครงสร้างรวม",
             "Pearson ≈ Spearman → ความสัมพันธ์เชิงเส้น",
             "ตัวแปรอินพุตสัมพันธ์กันเองต่ำ → ไม่มี multicollinearity"],
            kicker="ส่วนที่ 3 · numeric × numeric", accent=SAGE)

# ============================================================ relationship charts
chart_slide("ปัจจัยตัวเลข แยกตามสถานะ Has_KT", "07_violin_by_haskt.png",
            ["#รูปทรงการกระจายแยกกลุ่ม",
             "กลุ่มติด KT: SGPA / เข้าเรียน / ชั่วโมงอ่าน ต่ำกว่าชัดเจน",
             "กลุ่มติด KT: โซเชียล / เวลาเดินทาง สูงกว่า",
             "ความต่างของมัธยฐานสองกลุ่มเห็นชัดทุกปัจจัย",
             "การนอนต่างกันเล็กน้อย"],
            kicker="ส่วนที่ 3 · numeric × categorical", accent=SAGE)
chart_slide("ค่าเฉลี่ยจำนวน KT แยกตามตัวแปรกลุ่ม", "08_meanKT_by_category.png",
            ["#ปัจจัยกลุ่มที่แยกได้ชัด",
             "Stress_Level: เฉลี่ย 0.6 → 1.9 → 3.5 วิชา",
             "Commute_Mode: Local Train แย่กว่าโหมดอื่น",
             "มี Peer_Learning / Mentorship = KT น้อยกว่า",
             "#แถบ error = ช่วงเชื่อมั่น 95%"],
            kicker="ส่วนที่ 3 · categorical × target", accent=SAGE)
chart_slide("ความสัมพันธ์ กลุ่ม × กลุ่ม (Cramér's V)", "09_cramers_v.png",
            ["#Cramér's V (0–1)",
             "Stress_Level × Has_KT: สูงเด่นชัด (~0.56)",
             "Peer_Learning / Mentorship × Has_KT: อ่อน (~0.12)",
             "#ข้อสังเกต",
             "demographic (เพศ สาขา ชั้นปี ฐานะ) เป็นอิสระต่อกัน",
             "ชุดข้อมูลออกแบบให้ตัวแปรพื้นฐานไม่พันกัน"],
            kicker="ส่วนที่ 3 · categorical × categorical", accent=SAGE)
chart_slide("Stress_Level × Has_KT — ความสัมพันธ์ที่แรงที่สุด", "10_stress_vs_haskt.png",
            ["#สัดส่วนติด KT ตามระดับความเครียด",
             "Low → 50% · Medium → 98% · High → 100%",
             "ความสัมพันธ์แรงมาก (Cramér's V = 0.56)",
             "#ข้อควรระวัง",
             "ความเครียดอาจเป็น \"ผล\" ของการติด KT ด้วย",
             "→ อย่าใช้เป็น predictor แบบไร้เงื่อนไข"],
            kicker="ส่วนที่ 3 · relationship", accent=SAGE)
chart_slide("Pairplot ปัจจัยตัวเลขหลัก (สี = Has_KT)", "11_pairplot.png",
            ["#ภาพหลายมิติ",
             "กลุ่มสองสี (ติด / ไม่ติด KT) เหลื่อมกันมาก",
             "ไม่มีคู่ตัวแปรใดแยกกลุ่มได้ขาด",
             "#นัยยะ",
             "ต้องใช้หลายปัจจัยร่วมกันในการทำนาย",
             "เหมาะกับโมเดลที่รวมหลาย feature",
             "(เช่น logistic regression / decision tree)"],
            kicker="ส่วนที่ 3 · multivariate", accent=SAGE)

# ============================================================ key findings
def findings_slide(title, groups):
    s = base_slide(); footer(s)
    head(s, title, kicker="สรุป · Key Findings", accent=BLUE)
    y = Inches(2.15)
    for gcolor, gname, items in groups:
        card_h = Inches(0.78) + Inches(0.44) * len(items)
        _rect(s, Inches(0.72), y, Inches(11.9), card_h, tint(gcolor), round_=True, radius=0.05)
        _rect(s, Inches(0.72), y, Pt(5), card_h, gcolor, round_=True, radius=0.5)
        _text(s, Inches(1.08), y + Inches(0.18), Inches(11.2), Inches(0.5),
              [[(gname, dict(size=17, bold=True, color=gcolor, name=HEAD))]])
        _text(s, Inches(1.08), y + Inches(0.72), Inches(11.0), card_h - Inches(0.7),
              [[("•  " + it, dict(size=14, color=BODY, space_after=4))] for it in items],
              line_spacing=1.2)
        y = y + card_h + Inches(0.28)
    return s

findings_slide("สรุปข้อค้นพบ (1/2)", [
    (BLUE, "โครงสร้าง & คุณภาพข้อมูล",
     ["10,000 × 21 · ไม่มีแถว/ID ซ้ำ · ค่าหาย ~2% เฉพาะ 4 คอลัมน์ตัวเลข กระจายแบบสุ่ม",
      "Has_KT, Cleared_in_First_Attempt, (Number_of_KTs>0) = ข้อมูลชุดเดียวกัน → ระวัง target leakage"]),
    (RUST, "สถิติเชิงพรรณนา",
     ["ตัวแปรตัวเลขแจกแจงสมมาตร (|skew| < 0.8) · kurtosis ติดลบ (หางบาง) คล้ายสุ่มจาก uniform/normal",
      "เป้าหมายไม่สมดุล: 74% ติด KT · เฉลี่ย 1.54 วิชา/คน · สูงสุด 7"]),
])
findings_slide("สรุปข้อค้นพบ (2/2)", [
    (SAGE, "ความสัมพันธ์",
     ["ลด KT: SGPA เดิม (r≈−0.35), ชั่วโมงอ่านหนังสือ (−0.30), การเข้าเรียน (−0.30)",
      "เพิ่ม KT: ชั่วโมงโซเชียล (r≈+0.18), เวลาเดินทาง (+0.11); การนอน/อายุ แทบไม่มีผล",
      "ปัจจัยกลุ่มที่แยกความเสี่ยงชัด: Stress_Level (0.6 → 3.5) รองลงมา Commute_Mode, Peer_Learning, Mentorship",
      "ไม่มีผล: เพศ, สาขา, ชั้นปี, ฐานะการเงิน, กิจกรรมนอกหลักสูตร",
      "ตัวแปรอินพุตสัมพันธ์กันเองต่ำ → ทุกตัวให้ข้อมูลเสริมกัน ไม่มี multicollinearity"]),
    (GOLD, "นัยเชิงปฏิบัติ",
     ["กลุ่ม \"เครียดสูง + เข้าเรียนน้อย + โซเชียลเยอะ + SGPA ต่ำ\" คือเป้าหมายของระบบ early-warning"]),
])

# ============================================================ 25 · closing
s = base_slide(framed=False)
_rect(s, 0, 0, SW, SH, INK)
for i, c in enumerate([BLUE, RUST, SAGE]):
    _rect(s, Inches(0.95) + i * Inches(0.5), Inches(2.55), Inches(0.32), Inches(0.32), c, round_=True, radius=0.3)
_rect(s, Inches(0.95), Inches(3.15), Inches(1.7), Pt(4), RUST, round_=True, radius=0.5)
_text(s, Inches(0.95), Inches(3.4), Inches(11.5), Inches(1.2),
      [[("ขอบคุณครับ", dict(size=40, bold=True, color=WHITE, name=HEAD))]])
_text(s, Inches(0.97), Inches(4.7), Inches(11.5), Inches(1.4),
      [[("ชื่อ–สกุล: ", dict(size=18, color=RGBColor(0xCF, 0xC7, 0xBA))),
        (AUTHOR, dict(size=18, bold=True, color=WHITE))],
       [("รหัสนักศึกษา: ", dict(size=18, color=RGBColor(0xCF, 0xC7, 0xBA))),
        (SID, dict(size=18, bold=True, color=WHITE))]], line_spacing=1.4)

out = "Mumbai_KT_EDA_Presentation.pptx"
prs.save(out)
print("saved", out, "|", len(prs.slides._sldIdLst), "slides")
